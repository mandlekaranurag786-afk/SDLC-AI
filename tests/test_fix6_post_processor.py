import tempfile
from pathlib import Path
import pytest
from docx import Document
from post_processor import scan_placeholders, render_markdown_table, add_cover_warning_banner
from document_export import build_docx

def test_placeholder_scanning():
    md = """# Section 1
The total cost is XX USD.
Completion date is TBD.

## Appendix: Open Items for Review
- [HIGH] Section 1: TBD item in appendix should be ignored.
"""
    issues = scan_placeholders(md, doc_name="Test Doc")
    assert len(issues) == 2
    assert issues[0]["severity"] == "HIGH"
    assert "XX" in issues[0]["recommendation"]
    assert "TBD" in issues[1]["recommendation"]


def test_markdown_table_conversion():
    document = Document()
    table_md = [
        "| Feature | Priority | Effort |",
        "| --- | --- | --- |",
        "| SSO Login | High | 5d |",
        "| HR Dashboard | Medium | 8d |"
    ]
    render_markdown_table(document, table_md)
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].cells[0].text == "Feature"
    assert table.rows[1].cells[0].text == "SSO Login"
    
    # Verify repeating header xml element <w:tblHeader/>
    header_trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = header_trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader")
    assert tblHeader is not None


def test_cover_warning_banner():
    document = Document()
    report = {
        "issues": [
            {"severity": "HIGH", "section": "BRD", "affected_requirement": "Req 1", "recommendation": "Missing integration"}
        ]
    }
    add_cover_warning_banner(document, report)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "DRAFT — CONTAINS UNRESOLVED REQUIREMENTS ISSUES" in full_text


def test_build_docx_with_post_processing():
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = Path(tmpdir) / "test.docx"
        md_content = """# Business Requirements
| Component | Status |
| --- | --- |
| API Gateway | Active |
| Auth Service | XX |
"""
        report = {"issues": []}
        build_docx(
            title="Business Requirement Document",
            project_name="Test Project",
            markdown=md_content,
            output_path=out_path,
            traceability=[],
            include_traceability=False,
            validation_report=report
        )
        assert out_path.exists()
        doc = Document(out_path)
        assert len(doc.tables) >= 1
        # Assert placeholder XX was caught and added to report
        assert any("XX" in str(issue.get("recommendation")) for issue in report["issues"])
