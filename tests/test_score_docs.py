import pytest
from pathlib import Path
from docx import Document
import sys
import os

# Ensure tools and root are on sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))

from tools.score_docs import score_documents, parse_input_markdown, stem_word, format_scoreboard


def test_stem_word():
    assert stem_word("notifications") == "notification"
    assert stem_word("reminders") == "reminder"
    assert stem_word("provisioning") == "provision"
    assert stem_word("workflows") == "workflow"


def test_parse_input_markdown(tmp_path):
    md_file = tmp_path / "sample_input.md"
    md_file.write_text("""
# Problem Statement
Our organization needs a system.

# Q1. What additional functional features do you envision?
* Automated checklist
* Document upload

# Q2. Desired user workflows
HR:
* Create onboarding request
IT:
* Provision laptop

# Q3. What external systems should this application integrate with?
* HRMS
* Slack

# Q4. Constraints
* Technology: Python backend
* Budget: Low cost

# Q5. Success Metrics
* Reduction in onboarding time
* High satisfaction score
""", encoding="utf-8")

    categories = parse_input_markdown(md_file.read_text(encoding="utf-8"))
    assert len(categories["features"]) == 2
    assert len(categories["workflow steps"]) == 2
    assert len(categories["integrations"]) == 2
    assert len(categories["constraints"]) == 2
    assert len(categories["metrics"]) == 2


def test_score_documents_with_seeded_docx(tmp_path):
    input_file = tmp_path / "test_input.md"
    input_file.write_text("""
# Problem Statement
We need an automated employee onboarding system.

# Q1. What additional functional features do you envision?
* Automated checklist and task assignment
* Document upload and verification

# Q2. Desired user workflows
HR:
* Create new employee onboarding request
IT:
* Provision laptop and email

# Q3. What external systems should this application integrate with?
* HRMS System
* Slack

# Q4. Constraints
* Technology: Python FastAPI
* Compliance: Data privacy

# Q5. Success Metrics
* Reduction in average onboarding time
* Task completion SLA
""", encoding="utf-8")

    # 1. BRD Document
    brd_doc = Document()
    brd_doc.core_properties.title = "Business Requirement Document"
    brd_doc.core_properties.subject = "We Spend Significant Time Onboarding"  # Insane title: leading stopword
    brd_doc.add_paragraph("Business Requirement Document")
    brd_doc.add_paragraph("IT must provision assets within 24 hours.")  # BRD timing: 24 hours
    brd_doc.add_paragraph("Automated reminders every 24 hours.")
    brd_doc.add_paragraph("Escalations triggered after 24 hours.")
    # Leaked banned phrase and validator leak in body
    brd_doc.add_paragraph("This shall be clarified in subsequent phases.")
    brd_doc.add_paragraph("[HIGH] Section 5.1: Missing requirement detail")
    brd_path = tmp_path / "BRD.docx"
    brd_doc.save(str(brd_path))

    # 2. FRD Document
    frd_doc = Document()
    frd_doc.core_properties.title = "Functional Requirement Document"
    frd_doc.core_properties.subject = "Smart Onboarding System"  # Sane title
    frd_doc.add_heading("FR-001: Automated Checklist and Task Assignment", level=1)
    frd_doc.add_paragraph("Requirement ID: FR-001")
    frd_doc.add_paragraph("Description: Automated checklist and task assignment for onboarding.")
    frd_doc.add_paragraph("IT shall provision laptop and email within 48 hours.")  # FRD timing: 48 hours -> CONFLICT with BRD
    frd_doc.add_paragraph("Traceability: Maps to BRD Section: 5.1")
    # Tags on non-requirement values:
    frd_doc.add_paragraph("Requirement ID: FR-001 [Assumption/Proposal: confirm]")
    frd_doc.add_paragraph("1 [Proposed default — confirm]. Step one starts.")
    # Untagged number with unit:
    frd_doc.add_paragraph("The maximum file size supported is 50 MB.")
    frd_path = tmp_path / "FRD.docx"
    frd_doc.save(str(frd_path))

    # 3. Epics Document
    epics_doc = Document()
    epics_doc.core_properties.title = "Epics & User Stories"
    epics_doc.core_properties.subject = "Smart Onboarding System"
    # Duplicate story ID US-001:
    epics_doc.add_heading("US-001: Create Request", level=2)
    epics_doc.add_paragraph("As an HR user I want to create request.")
    epics_doc.add_paragraph("Traceability: FR-001")

    epics_doc.add_heading("US-001: Duplicate Header", level=2)  # Duplicate!
    epics_doc.add_paragraph("As an admin I want to duplicate.")
    # Story without FR ID:
    epics_doc.add_heading("US-002: Provision Laptop", level=2)
    epics_doc.add_paragraph("As an IT engineer I want to provision laptop.")
    # Notice: no FR-xxx cited!

    epics_path = tmp_path / "Epics_Stories.docx"
    epics_doc.save(str(epics_path))

    # 4. Estimation Document
    est_doc = Document()
    est_doc.core_properties.title = "Project Estimation Report"
    est_doc.core_properties.subject = "Smart Onboarding System"
    # Reference dangling FR-999 and FR-100:
    t = est_doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Sprint 1"
    t.rows[0].cells[1].text = "FR-001"
    t.rows[1].cells[0].text = "Sprint 2"
    t.rows[1].cells[1].text = "FR-999 to FR-100"  # Dangling!
    est_path = tmp_path / "Project_Estimation.docx"
    est_doc.save(str(est_path))

    scores = score_documents(
        input_path=input_file,
        brd_path=brd_path,
        frd_path=frd_path,
        epics_path=epics_path,
        estimation_path=est_path,
    )

    # Validate computed metrics
    assert scores["duplicate_story_ids"] == 1  # US-001 repeated
    assert scores["stories_without_fr"] >= 1   # US-002 has no FR link
    assert scores["dangling_fr_count"] >= 1    # FR-999 / FR-100 not in FRD (only FR-001 defined)
    assert scores["untagged_numbers_with_units"] >= 1  # 50 MB is not in input and untagged
    assert scores["tags_on_non_requirement_values"] >= 2  # FR-001 [tag] and 1 [tag]
    assert scores["brd_frd_timing_conflicts"] >= 1  # Provisioning: 24 hours vs 48 hours
    assert scores["banned_phrases_count"] >= 2  # "shall be clarified" and "[HIGH]"
    assert scores["title_sanity_issues"] >= 1   # "We Spend Significant Time Onboarding" has stopword

    output_str = format_scoreboard(scores)
    assert "SDLC DOCUMENT GENERATION SCOREBOARD" in output_str
