import re
from typing import Any
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PLACEHOLDER_PATTERNS = [
    (r"\bXX\b", "Unresolved placeholder 'XX' found"),
    (r"\bTBD\b", "Unresolved placeholder 'TBD' found"),
    (r"\[insert\b", "Unresolved placeholder '[insert...]' found"),
    (r"\<insert\b", "Unresolved placeholder '<insert...>' found"),
    (r"\bTODO\b", "Unresolved placeholder 'TODO' found"),
]

CHATTY_CLOSING_PATTERNS = [
    re.compile(r"\n*(?:This concludes|End of document|This concludes the|End of section)[\s\S]*$", re.IGNORECASE),
    re.compile(r"\n*(?:Note:\s*ensure that all items|Note:\s*shall be clarified|Note:\s*must be defined)[\s\S]*$", re.IGNORECASE),
]


def strip_chatty_closings(markdown_content: str) -> str:
    """
    Strips trailing chatty LLM sign-offs, closing summaries, and meta-notes from document text.
    """
    if not markdown_content:
        return markdown_content

    cleaned = markdown_content.strip()
    for pattern in CHATTY_CLOSING_PATTERNS:
        cleaned = pattern.sub("", cleaned)
    return cleaned.strip()

def scan_placeholders(markdown_content: str, doc_name: str = "Document") -> list[dict[str, Any]]:
    """
    Scans markdown text for placeholder tokens (XX, TBD, [insert, <insert, TODO).
    Ignores lines inside 'Appendix: Open Items'.
    Allows intentional 'Target: TBD / to be defined' only under 'Project Objectives' and 'Success Criteria'.
    Returns list of validation issue dicts with severity HIGH.
    """
    from feature_flags import ENABLE_VALIDATOR_FP_FIXES

    issues = []
    in_open_items = False
    current_heading = ""

    for line_idx, line in enumerate(markdown_content.splitlines(), start=1):
        stripped = line.strip()
        if stripped.startswith("#"):
            current_heading = stripped
            if "open items" in stripped.lower():
                in_open_items = True
                continue
            else:
                in_open_items = False

        if in_open_items:
            continue

        if ENABLE_VALIDATOR_FP_FIXES:
            is_objectives_or_success = any(h in current_heading.lower() for h in ["objectives", "success criteria"])
            if is_objectives_or_success and re.search(r"Target:\s*(?:TBD|to be defined)", line, re.IGNORECASE):
                continue

        for pattern, msg in PLACEHOLDER_PATTERNS:
            if re.search(pattern, line, flags=re.IGNORECASE):
                issues.append({
                    "type": "placeholder_leakage",
                    "severity": "HIGH",
                    "section": doc_name,
                    "affected_requirement": f"Line {line_idx}",
                    "recommendation": f"{msg}: '{stripped[:60]}...'"
                })
                break

    return issues


def add_cover_warning_banner(document: Document, validation_report: dict[str, Any] | None) -> None:
    """
    If validation_report contains any HIGH or CRITICAL issues,
    adds a prominent warning banner on the cover page.
    """
    if not validation_report:
        return

    issues = validation_report.get("issues") or []
    has_high_or_critical = any(
        str(issue.get("severity", "")).upper() in {"HIGH", "CRITICAL"}
        for issue in issues if isinstance(issue, dict)
    )

    if not has_high_or_critical:
        return

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(12)
    
    run = paragraph.add_run("⚠️ DRAFT — CONTAINS UNRESOLVED REQUIREMENTS ISSUES")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(220, 38, 38)  # Red banner text


def render_markdown_table(document: Document, table_lines: list[str]) -> None:
    """
    Parses contiguous markdown table lines (| col1 | col2 |) into a native python-docx Table.
    Sets Table Grid style and appends <w:tblHeader/> to header row trPr so headers repeat across page breaks.
    """
    if not table_lines:
        return

    cleaned_rows = []
    for line in table_lines:
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            continue
        cells = [c.strip() for c in stripped.split("|")[1:-1]]
        # Skip alignment separator rows like |---|---|
        if all(re.match(r"^:?-+:?$", c) for c in cells):
            continue
        cleaned_rows.append(cells)

    if not cleaned_rows:
        return

    num_rows = len(cleaned_rows)
    num_cols = max(len(row) for row in cleaned_rows)

    table = document.add_table(rows=num_rows, cols=num_cols)
    table.style = "Table Grid"

    # Set <w:tblHeader/> on header row (row 0)
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement("w:tblHeader"))

    for r_idx, row_data in enumerate(cleaned_rows):
        row_cells = table.rows[r_idx].cells
        for c_idx, cell_text in enumerate(row_data):
            if c_idx < len(row_cells):
                cell_paragraph = row_cells[c_idx].paragraphs[0]
                cell_paragraph.text = cell_text
                if r_idx == 0:
                    for run in cell_paragraph.runs:
                        run.bold = True
